#!/usr/bin/env python3
"""Create the detailed Jetson voice-agent benchmark presentation."""

from __future__ import annotations

import json
import math
import statistics
import sys
from pathlib import Path

from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt


SKILL_ROOT = Path("/home/thangnv94/.codex/skills/friendly-project-slides")
sys.path.insert(0, str(SKILL_ROOT / "scripts"))

from slidekit import (  # noqa: E402
    BG,
    FONT,
    INK,
    LINE,
    MUTED,
    NAVY,
    ORANGE,
    ORANGE_PALE,
    PAPER,
    SOFT,
    TEAL,
    TEAL_DARK,
    TEAL_PALE,
    WHITE,
    add_footer,
    add_module,
    add_notes,
    add_text,
    add_title,
    blank_slide,
    new_presentation,
    rgb,
)


REPO_ROOT = Path(__file__).resolve().parents[1]
OUTPUT = REPO_ROOT / "docs" / "jetson-voice-agent-results-by-audio-length.pptx"
BENCHMARK_ROOT = REPO_ROOT / "outputs" / "benchmarks"

TOOL_SAMPLES = [
    ("01_1.39s_radio_source.wav", "1.39 s"),
    ("02_2.00s_fog_lights.wav", "2.00 s"),
    ("03_3.00s_seat_heating.wav", "3.00 s"),
    ("04_4.00s_ambient_green.wav", "4.00 s"),
    ("05_5.00s_ambient_rainbow.wav", "5.00 s"),
    ("06_6.00s_ambient_rainbow.wav", "6.00 s"),
    ("07_6.75s_next_station.wav", "6.75 s"),
    ("08_7.44s_previous_station.wav", "7.44 s"),
]

NON_TOOL_SAMPLES = [
    ("01_1.000s_news.wav", "1.0 s"),
    ("02_3.000s_news.wav", "3.0 s"),
    ("03_5.000s_news.wav", "5.0 s"),
    ("04_7.000s_news.wav", "7.0 s"),
    ("05_7.500s_news.wav", "7.5 s"),
]


def percentile(values: list[float], fraction: float) -> float:
    ordered = sorted(values)
    position = (len(ordered) - 1) * fraction
    lower = math.floor(position)
    upper = math.ceil(position)
    if lower == upper:
        return ordered[lower]
    weight = position - lower
    return ordered[lower] * (1 - weight) + ordered[upper] * weight


def metric_value(record: dict, metric: str) -> float:
    if metric == "ttft":
        return float(record["audio_ready_to_first_slm_token_ms"])
    if metric == "generation":
        return float(record["slm"]["timings"]["first_to_last_token_ms"])
    if metric == "tps":
        generation_s = metric_value(record, "generation") / 1000
        return (float(record["slm"]["output_tokens"]) - 1) / generation_s
    if metric == "last_slm":
        return float(record["audio_ready_to_last_slm_token_ms"])
    if metric == "cloud":
        return float(record["audio_ready_to_cloud_llm_ms"])
    if metric == "first_audio":
        return float(record["audio_ready_to_first_audio_byte_ms"])
    raise KeyError(metric)


def stat_cell(records: list[dict], metric: str) -> str:
    if not records:
        return "— / — / —"
    values = [metric_value(record, metric) for record in records]
    stats = (statistics.median(values), percentile(values, 0.90), max(values))
    if metric == "tps":
        return " / ".join(f"{value:.1f}" for value in stats)
    return " / ".join(f"{round(value):,}" for value in stats)


def benchmark_records(gpu_share: int, *, non_tool: bool) -> list[dict]:
    folder = "jetson_cloud_path_20260814" if non_tool else "jetson_full_system_20260814"
    valid_key = "valid_cloud_path" if non_tool else "valid_tool_path"
    paths = [BENCHMARK_ROOT / folder / f"p{gpu_share}" / "runs.jsonl"]
    if non_tool and gpu_share == 30:
        paths.append(
            BENCHMARK_ROOT
            / folder
            / "p30_missing_3s_7s"
            / "runs.jsonl"
        )
    records = [
        json.loads(line)
        for path in paths
        for line in path.read_text().splitlines()
        if line.strip()
    ]
    return [
        record
        for record in records
        if record.get("measured") and record.get(valid_key)
    ]


def benchmark_rows(gpu_share: int, *, non_tool: bool) -> list[list[str]]:
    records = benchmark_records(gpu_share, non_tool=non_tool)
    samples = NON_TOOL_SAMPLES if non_tool else TOOL_SAMPLES
    metrics = (
        ("ttft", "generation", "tps", "last_slm", "cloud", "first_audio")
        if non_tool
        else ("ttft", "generation", "tps", "last_slm", "first_audio")
    )
    rows = []
    for filename, label in samples:
        selected = [record for record in records if record["sample"] == filename]
        rows.append([label, *(stat_cell(selected, metric) for metric in metrics)])
    return rows


def aggregate_tool_rows() -> list[list[str]]:
    metrics = ("ttft", "generation", "tps", "last_slm", "first_audio")
    return [
        [
            "Full GPU\nMPS off" if gpu_share == 100 else f"MPS {gpu_share}%",
            *(stat_cell(benchmark_records(gpu_share, non_tool=False), metric) for metric in metrics),
        ]
        for gpu_share in (100, 70, 50, 30)
    ]


def aggregate_non_tool_rows() -> list[list[str]]:
    metrics = ("ttft", "generation", "tps", "last_slm", "cloud", "first_audio")
    return [
        [
            "Full GPU\nMPS off" if gpu_share == 100 else f"MPS {gpu_share}%",
            *(stat_cell(benchmark_records(gpu_share, non_tool=True), metric) for metric in metrics),
        ]
        for gpu_share in (100, 70, 50, 30)
    ]


TOOL_SUMMARY = {
    100: ("Full-GPU baseline keeps TTFT near 0.2 seconds", "MPS disabled · 114.9 tok/s aggregate decode · 568 ms to first cached audio"),
    70: ("70% MPS closely tracks full-thread latency", "108.6 tok/s aggregate decode · only +28 ms to first audio vs 100%"),
    50: ("50% MPS remains interactive at 101.5 tok/s", "632 ms from VAD-ready audio to first cached audio · +64 ms vs 100%"),
    30: ("30% MPS slows long structured calls", "54.8 tok/s aggregate decode · longest calls need about 1.49 s to finish"),
}


NON_TOOL_SUMMARY = {
    100: ("Cloud + TTS dominate the full-GPU baseline", "MPS disabled · 99.5 tok/s SLM decode · Gemini 2.13 s median"),
    70: ("70% MPS keeps local routing fast", "94.4 tok/s SLM decode · external Gemini, Wi-Fi, and TTS remain variable"),
    50: ("50% MPS preserves local routing", "88.8 tok/s SLM decode · post-VAD TTFT remains about 0.19–0.21 s"),
    30: ("30% MPS slows SLM routing", "local SLM decode falls to 45–64 tok/s by audio length"),
}


def set_cell_border(cell, color: str = LINE, width: int = 9500) -> None:
    """Apply a quiet grid border to a PowerPoint table cell."""
    tc_pr = cell._tc.get_or_add_tcPr()
    for edge in ("lnL", "lnR", "lnT", "lnB"):
        for old in tc_pr.findall(f"{{http://schemas.openxmlformats.org/drawingml/2006/main}}{edge}"):
            tc_pr.remove(old)
        line = OxmlElement(f"a:{edge}")
        line.set("w", str(width))
        line.set("cap", "flat")
        line.set("cmpd", "sng")
        line.set("algn", "ctr")
        fill = OxmlElement("a:solidFill")
        srgb = OxmlElement("a:srgbClr")
        srgb.set("val", color)
        fill.append(srgb)
        line.append(fill)
        dash = OxmlElement("a:prstDash")
        dash.set("val", "solid")
        line.append(dash)
        tc_pr.append(line)


def add_table(
    slide,
    headers: list[str],
    rows: list[list[str]],
    x: float,
    y: float,
    w: float,
    h: float,
    widths: list[float],
    *,
    header_size: float = 10.5,
    body_size: float = 13.0,
    highlight_rows: dict[int, str] | None = None,
    left_columns: set[int] | None = None,
):
    """Add a compact, editable PowerPoint table."""
    if len(headers) != len(widths):
        raise ValueError("headers and widths must have equal length")
    if abs(sum(widths) - w) > 0.02:
        raise ValueError("column widths must sum to the table width")

    shape = slide.shapes.add_table(
        len(rows) + 1,
        len(headers),
        Inches(x),
        Inches(y),
        Inches(w),
        Inches(h),
    )
    table = shape.table
    table.first_row = True
    table.horz_banding = False
    for col, col_width in zip(table.columns, widths):
        col.width = Inches(col_width)

    row_h = h / (len(rows) + 1)
    for row in table.rows:
        row.height = Inches(row_h)

    all_rows = [headers] + rows
    highlight_rows = highlight_rows or {}
    left_columns = left_columns or set()

    for r_idx, values in enumerate(all_rows):
        for c_idx, value in enumerate(values):
            cell = table.cell(r_idx, c_idx)
            fill_color = NAVY if r_idx == 0 else highlight_rows.get(r_idx - 1, PAPER)
            cell.fill.solid()
            cell.fill.fore_color.rgb = rgb(fill_color)
            set_cell_border(cell)
            cell.margin_left = Inches(0.08)
            cell.margin_right = Inches(0.08)
            cell.margin_top = Inches(0.05)
            cell.margin_bottom = Inches(0.05)

            tf = cell.text_frame
            tf.clear()
            tf.word_wrap = True
            tf.vertical_anchor = MSO_ANCHOR.MIDDLE
            for line_idx, line in enumerate(str(value).split("\n")):
                p = tf.paragraphs[0] if line_idx == 0 else tf.add_paragraph()
                p.alignment = PP_ALIGN.LEFT if c_idx in left_columns else PP_ALIGN.CENTER
                p.space_before = Pt(0)
                p.space_after = Pt(0)
                p.line_spacing = 1.0
                run = p.add_run()
                run.text = line
                run.font.name = FONT
                run.font.size = Pt(header_size if r_idx == 0 else body_size)
                run.font.bold = r_idx == 0 or c_idx == 0
                run.font.color.rgb = RGBColor.from_string(WHITE if r_idx == 0 else INK)
    return shape


def slide_how_to_read(prs) -> None:
    slide = blank_slide(prs, BG)
    add_title(
        slide,
        "How to read the report",
        "The reported clock starts after VAD",
        "VAD endpointing is excluded because the test used synthetic clean trailing silence.",
    )

    add_text(
        slide,
        "Post-VAD time to first audio = SLM TTFT + SLM generation + [cloud LLM] + response preparation",
        0.82,
        2.08,
        11.70,
        0.72,
        size=18,
        color=TEAL_DARK,
        bold=True,
        align=PP_ALIGN.CENTER,
        valign=MSO_ANCHOR.MIDDLE,
    )
    add_table(
        slide,
        ["Stage duration", "What it measures"],
        [
            ["Reported statistics", "Each result cell shows P50 / P90 / Max from left to right"],
            ["SLM TTFT", "Finalized audio → first meaningful SLM token"],
            ["SLM generation", "First → last SLM token; TPS = (output tokens − 1) / this duration"],
            ["Cloud LLM", "Only for non_tool; SLM completion → complete Gemini response"],
            ["Response preparation", "Cloud TTS or cached reply → first response-audio byte"],
        ],
        0.82,
        2.95,
        11.70,
        3.45,
        [2.75, 8.95],
        header_size=11.5,
        body_size=11.3,
        left_columns={0, 1},
    )
    add_text(
        slide,
        "“Audio → …” means VAD-ready audio → milestone. It does not include speaking time or VAD endpointing.",
        0.84,
        6.58,
        11.60,
        0.28,
        size=10,
        color=TEAL_DARK,
        bold=True,
        align=PP_ALIGN.CENTER,
    )
    add_footer(slide, 1, "Measured on Jetson AGX Orin · 14 Aug 2026 · VAD endpoint latency excluded")
    add_notes(
        slide,
        "All displayed latency begins when VAD has finalized the waveform. "
        "The previous endpoint-delay result was removed because the runner appended synthetic zero-valued silence, which does not represent a naturally noisy robot environment. "
        "The equation uses non-overlapping post-VAD stage durations; cumulative milestone columns should not be summed. Audio-to-last-audio is intentionally omitted.",
    )


def slide_tool_aggregate(prs, number: int) -> None:
    slide = blank_slide(prs, BG)
    add_title(
        slide,
        "TOOL-CALL SUMMARY · ALL AUDIO LENGTHS",
        "50% MPS keeps tool calls interactive",
        "64 measured turns per scheduling condition · every result cell is P50 / P90 / Max",
    )
    add_table(
        slide,
        [
            "GPU scheduling\ncondition",
            "Audio → 1st SLM\nP50 / P90 / Max",
            "1st → last SLM\nP50 / P90 / Max",
            "SLM TPS\nP50 / P90 / Max",
            "Audio → last SLM\nP50 / P90 / Max",
            "Audio → 1st audio\nP50 / P90 / Max",
        ],
        aggregate_tool_rows(),
        0.72,
        2.18,
        11.89,
        3.42,
        [1.10, 2.15, 2.15, 1.65, 2.20, 2.64],
        header_size=9.3,
        body_size=9.8,
    )
    add_text(
        slide,
        "At the 50% MPS cap, P50 is 219 ms to first SLM token, 595 ms to the completed tool call, and 632 ms to first cached audio.",
        0.90,
        5.90,
        11.50,
        0.54,
        size=13.2,
        color=TEAL_DARK,
        bold=True,
        align=PP_ALIGN.CENTER,
        valign=MSO_ANCHOR.MIDDLE,
    )
    add_footer(
        slide,
        number,
        "All 8 audio lengths pooled · 64 warm measured turns per condition · VAD and last audio excluded",
    )
    add_notes(
        slide,
        "This is the aggregate tool-call view requested for a single overall distribution across every tested audio length. "
        "Each cell reads P50, P90, and maximum from left to right. The 50% MPS active-thread cap remains close to full-thread execution, while 30% materially increases generation and completion latency. "
        "These are project measurements from the controlled tool-call corpus; the detailed per-length slides follow.",
    )


def slide_tool_results(prs, gpu_share: int, number: int) -> None:
    slide = blank_slide(prs, BG)
    title, subtitle = TOOL_SUMMARY[gpu_share]
    add_title(
        slide,
        "Tool-call results · full GPU · MPS off" if gpu_share == 100 else f"Tool-call results · {gpu_share}% MPS cap",
        title,
        f"{subtitle} · each cell shows P50 / P90 / Max",
    )

    add_table(
        slide,
        [
            "Audio\nlength",
            "Audio → 1st SLM\nP50 / P90 / Max",
            "1st → last SLM\nP50 / P90 / Max",
            "SLM TPS\nP50 / P90 / Max",
            "Audio → last SLM\nP50 / P90 / Max",
            "Audio → 1st audio\nP50 / P90 / Max",
        ],
        benchmark_rows(gpu_share, non_tool=False),
        0.72,
        1.94,
        11.89,
        4.78,
        [1.10, 2.10, 2.10, 1.60, 2.10, 2.89],
        header_size=8.5,
        body_size=8.9,
    )
    add_footer(
        slide,
        number,
        "Statistics: P50 / P90 / Max · 8 runs per audio · TPS excludes TTFT · VAD and last audio excluded",
    )
    add_notes(
        slide,
        ("This slide shows the controlled tool-call path on the direct full-GPU baseline with MPS disabled. " if gpu_share == 100 else f"This slide shows the controlled tool-call path at a {gpu_share}% CUDA MPS active-thread cap. ")
        + "Audio length remains visible rather than being aggregated away. "
        f"The aggregate decode rate is {TOOL_SUMMARY[gpu_share][1].split(' tok/s')[0]} tokens per second. "
        "Longer input audio does not automatically increase TTFT; the structured output length controls most of the first-to-last-token gap. "
        "Audio-to-last-audio is omitted because it is derived from response WAV duration rather than acoustic loopback.",
    )


def slide_non_tool_results(prs, gpu_share: int, number: int) -> None:
    slide = blank_slide(prs, BG)
    title, subtitle = NON_TOOL_SUMMARY[gpu_share]
    if gpu_share in (100, 70):
        sample_note = "6 uncached runs per audio length"
    elif gpu_share == 50:
        sample_note = "runs by length: 4 / 4 / 3 / 3 / 4"
    else:
        sample_note = "2 uncached runs per audio length"

    add_title(
        slide,
        "Non-tool results · full GPU · MPS off" if gpu_share == 100 else f"Non-tool results · {gpu_share}% MPS cap",
        title,
        f"{subtitle} · {sample_note}",
    )

    add_table(
        slide,
        [
            "Audio\nlength",
            "Audio → 1st SLM\nP50 / P90 / Max",
            "1st → last SLM\nP50 / P90 / Max",
            "SLM TPS\nP50 / P90 / Max",
            "Audio → last SLM\nP50 / P90 / Max",
            "Audio → cloud*\nP50 / P90 / Max",
            "Audio → 1st audio\nP50 / P90 / Max",
        ],
        benchmark_rows(gpu_share, non_tool=True),
        0.72,
        2.02,
        11.89,
        4.15,
        [1.00, 1.75, 1.75, 1.35, 1.75, 2.00, 2.29],
        header_size=8.1,
        body_size=8.2,
    )
    add_text(
        slide,
        "Cloud / network / laptop TTS are external; compare SLM TTFT, generation, and TPS across local scheduling conditions—not cloud totals.",
        0.84,
        6.34,
        11.60,
        0.32,
        size=10.2,
        color=TEAL_DARK,
        bold=True,
        align=PP_ALIGN.CENTER,
    )
    footer = (
        "Statistics: P50 / P90 / Max · *Gemini cloud completion, not observed last-token time · "
        "VAD/last audio excluded"
    )
    if gpu_share == 30:
        footer += " · P90 low-confidence (n=2)"
    add_footer(
        slide,
        number,
        footer,
    )
    add_notes(
        slide,
        ("This slide shows the non-tool path on the direct full-GPU baseline with MPS disabled. " if gpu_share == 100 else f"This slide shows the non-tool path at a {gpu_share}% CUDA MPS active-thread cap. ")
        + "The local SLM timings can be compared across local scheduling conditions, but Gemini, Wi-Fi, and laptop OmniVoice timings cannot be controlled by them. "
        "Gemini returned a complete non-streamed response, so the cloud milestone is response completion rather than a separately observed last token. "
        "P90 is linearly interpolated; the 30% non-tool P90 has only two runs and is descriptive, not a stable tail estimate.",
    )


def slide_non_tool_aggregate(prs, number: int) -> None:
    slide = blank_slide(prs, BG)
    add_title(
        slide,
        "NON-TOOL SUMMARY · ALL AVAILABLE RUNS",
        "Cloud and TTS dominate the non-tool path",
        "Runs: full GPU=30 · MPS70=30 · MPS50=18 · MPS30=10 · every cell is P50 / P90 / Max",
    )
    add_table(
        slide,
        [
            "GPU scheduling\ncondition",
            "Audio → 1st SLM\nP50 / P90 / Max",
            "1st → last SLM\nP50 / P90 / Max",
            "SLM TPS\nP50 / P90 / Max",
            "Audio → last SLM\nP50 / P90 / Max",
            "Audio → cloud*\nP50 / P90 / Max",
            "Audio → 1st audio\nP50 / P90 / Max",
        ],
        aggregate_non_tool_rows(),
        0.72,
        2.18,
        11.89,
        3.42,
        [1.00, 1.60, 1.60, 1.30, 1.60, 2.00, 2.79],
        header_size=8.1,
        body_size=8.1,
    )
    add_text(
        slide,
        "On the full-GPU baseline (MPS off), SLM routing completes at 667 ms P50; dynamic cloud speech starts at 5.24 s P50.",
        0.90,
        5.90,
        11.50,
        0.54,
        size=13.2,
        color=TEAL_DARK,
        bold=True,
        align=PP_ALIGN.CENTER,
        valign=MSO_ANCHOR.MIDDLE,
    )
    add_footer(
        slide,
        number,
        "*Gemini response completion · unequal repetitions by condition · external cloud/network/TTS latency · VAD/last audio excluded",
    )
    add_notes(
        slide,
        "This is the aggregate non-tool view across every available valid cloud-path run. Each cell reads P50, P90, and maximum from left to right. "
        "The local SLM decision remains much faster than the complete dynamic response path; Gemini, Wi-Fi, and laptop OmniVoice dominate the time to first audio. "
        "Do not rank MPS caps using the cloud columns: p50 has fewer repetitions, and external latency is uncontrolled. The 30% condition now covers all five audio lengths but still has only two repetitions per length.",
    )


def slide_deployment_summary(prs, number: int) -> None:
    slide = blank_slide(prs, BG)
    add_title(
        slide,
        "DEPLOYMENT SUMMARY",
        "Edge SLM is an option—not a requirement",
        "Choose placement from product constraints, not from architecture alone.",
    )
    add_module(
        slide,
        "Cloud-first is viable",
        "The complete voice pipeline—including the SLM—can run in the cloud. This is the simplest path to deploy and update.",
        0.82,
        2.28,
        3.73,
        2.30,
        accent=TEAL,
    )
    add_module(
        slide,
        "Edge SLM buys locality",
        "Local parsing can support low-latency tool calls, offline operation and tighter control of speech data.",
        4.80,
        2.28,
        3.73,
        2.30,
        accent=TEAL_DARK,
    )
    add_module(
        slide,
        "Edge costs engineering",
        "Jetson adds model/runtime compatibility, memory and MPS tuning, monitoring, thermal validation and fleet updates.",
        8.78,
        2.28,
        3.73,
        2.30,
        fill=ORANGE_PALE,
        accent=ORANGE,
    )
    add_text(
        slide,
        "Default to cloud; move the SLM to edge only when latency, connectivity, privacy or autonomy justify the extra deployment effort.",
        1.08,
        5.28,
        11.18,
        0.86,
        size=17,
        color=TEAL_DARK,
        bold=True,
        align=PP_ALIGN.CENTER,
        valign=MSO_ANCHOR.MIDDLE,
    )
    add_footer(
        slide,
        number,
        "Architectural takeaway from the measured Jetson deployment; placement remains a product decision",
    )
    add_notes(
        slide,
        "The whole pipeline can be hosted in the cloud, including the speech-language model. "
        "Running the SLM on Jetson is therefore a deliberate product trade-off, not an architectural requirement. "
        "Edge placement can improve locality, offline behavior and the bounded tool-call path, but it requires additional compatibility, resource-sharing, thermal and fleet-management work. "
        "The benchmark establishes that the edge path is feasible; it does not claim edge placement is always preferable.",
    )


def build(output: Path = OUTPUT) -> Path:
    prs = new_presentation(
        title="Jetson Voice-Agent Benchmark Results",
        author="Speech-to-Action project",
    )
    slide_how_to_read(prs)
    slide_tool_aggregate(prs, 2)
    for number, gpu_share in enumerate((100, 70, 50, 30), start=3):
        slide_tool_results(prs, gpu_share, number)
    slide_non_tool_aggregate(prs, 7)
    for number, gpu_share in enumerate((100, 70, 50, 30), start=8):
        slide_non_tool_results(prs, gpu_share, number)
    slide_deployment_summary(prs, 12)
    output.parent.mkdir(parents=True, exist_ok=True)
    prs.save(output)
    return output


if __name__ == "__main__":
    print(build())
